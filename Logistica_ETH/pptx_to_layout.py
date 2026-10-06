"""Read an ETH screen PowerPoint (one slide) and dump its layout as JSON.

The PowerPoint is the source of truth for the screen design: positions, text, fonts,
colours and pictures are read from it, so the video always matches the edited slide.
Coordinates are returned in screen pixels (the slides are 19.2 in wide for 1920 px,
i.e. 100 px per inch). Pictures are written next to the JSON.

Usage: python3 pptx_to_layout.py slide.pptx out_dir
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
EMU_PER_PX = 9144  # 914400 EMU per inch / 100 px per inch


def px(emu: int | None) -> float:
    return round((emu or 0) / EMU_PER_PX, 2)


def colour(node) -> tuple[str, float]:
    """Return (hex, alpha 0-1) for the first solidFill below node."""
    fill = node.find(".//a:solidFill", NS) if node is not None else None
    if fill is None:
        return "", 1.0
    clr = fill[0]
    value = clr.get("val", "FFFFFF")
    if clr.tag.endswith("schemeClr"):  # bg1 / tx1 etc: only white is used in these slides
        value = "FFFFFF" if value in ("bg1", "lt1") else "000000"
    alpha = clr.find("a:alpha", NS)
    return value.upper(), (int(alpha.get("val")) / 100000 if alpha is not None else 1.0)


def paragraphs(shape) -> list[dict]:
    out = []
    for p in shape.text_frame.paragraphs:
        ppr = p._p.find("a:pPr", NS)
        ln = ppr.find("a:lnSpc/a:spcPct", NS) if ppr is not None else None
        runs = []
        for r in p.runs:
            rpr = r._r.find("a:rPr", NS)
            hexv, _ = colour(rpr)
            latin = rpr.find("a:latin", NS) if rpr is not None else None
            runs.append({
                "t": r.text,
                "sz": round(int(rpr.get("sz", "1800")) / 100 / 72 * 100, 2),
                "b": rpr.get("b") == "1",
                "color": hexv or "FFFFFF",
                "face": latin.get("typeface") if latin is not None else "Arial",
                "spc": round(int(rpr.get("spc", "0")) / 100 / 72 * 100, 2),
            })
        if runs:
            out.append({
                "algn": (ppr.get("algn") if ppr is not None else None) or "l",
                "ln": int(ln.get("val")) / 100000 if ln is not None else 1.0,
                "runs": runs,
            })
    return out


def main(src: str, out_dir: str) -> None:
    out = Path(out_dir)
    (out / "media").mkdir(parents=True, exist_ok=True)
    prs = Presentation(src)
    slide = prs.slides[0]
    layout = {"w": round(px(prs.slide_width)), "h": round(px(prs.slide_height)), "items": []}
    for i, sh in enumerate(slide.shapes):
        item = {"i": i, "name": sh.name, "x": px(sh.left), "y": px(sh.top), "w": px(sh.width), "h": px(sh.height)}
        if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
            blob = sh.image
            fname = f"img{i}.{blob.ext}"
            (out / "media" / fname).write_bytes(blob.blob)
            item.update(type="pic", media=f"media/{fname}", pw=blob.size[0], ph=blob.size[1])
        elif sh.has_text_frame and sh.text_frame.text.strip():
            item.update(type="text", paras=paragraphs(sh))
        else:
            sppr = sh._element.find(".//{*}spPr")
            hexv, alpha = colour(sppr)
            geom = sppr.find("a:prstGeom", NS)
            adj = sppr.find(".//a:gd", NS)
            item.update(type="shape", prst=geom.get("prst") if geom is not None else "rect",
                        fill=hexv or "FFFFFF", alpha=alpha,
                        adj=int(adj.get("fmla").split()[1]) / 100000 if adj is not None else 0.0)
        layout["items"].append(item)
    (out / "layout.json").write_text(json.dumps(layout, indent=1), encoding="utf-8")
    print(f"{src}: {len(layout['items'])} items, {layout['w']}x{layout['h']}")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
