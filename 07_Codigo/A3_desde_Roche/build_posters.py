"""OILS26 A3 posters: teaser + agenda.

Builds each poster as a native PowerPoint file (every word is a live text box,
every logo a separate picture), then has PowerPoint itself export the PDF, so the
editable file and the print PDF are the same document.

    python3.12 make_art.py        # background artwork + QR (only when those change)
    python3.12 build_posters.py   # PPTX -> PDF (PowerPoint) -> PNG previews

All geometry is in millimetres on a 297 x 420 mm page.
Fonts: Eurostile (OILS identity) and Aptos — both ship with Microsoft Office.
"""
import copy
import os
import subprocess
import sys

from lxml import etree
from PIL import ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.package import Part
from pptx.opc.packuri import PackURI
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.join(HERE, "assets")
ROOT = os.path.abspath(os.path.join(HERE, "..", ".."))
OUTDIR = os.path.join(ROOT, "05_Versiones_anteriores", "2026-10_A3_rediseno_alt")
PREVIEW = os.path.join(OUTDIR, "previews")
FONTDIR = "/Applications/Microsoft PowerPoint.app/Contents/Resources/DFonts"

PAGE_W, PAGE_H = 297.0, 420.0
M = 20.0                      # outer margin
CW = PAGE_W - 2 * M           # content width

# ------------------------------------------------------------------ palette
NAVY_DEEP = "061433"
NAVY = "0B3E88"               # OILS logo navy
BLUE = "4B8CCD"               # OILS logo light blue
SKY = "8CCBF2"                # date / accent on navy
WHITE = "FFFFFF"
SOFT = "C9D6EC"               # secondary text on navy
MUTED = "93A7C9"              # tertiary text on navy
RULE = "3A6DAE"                # hairlines on navy (solid, so they stay vector in the PDF)
RULE_SOFT = "24487F"
INK = "0B2554"                # text on the white sponsor band

# ------------------------------------------------------------------ type
FONT_FILES = {
    ("Eurostile", False): "Eurostile.ttf",
    ("Eurostile", True): "Eurostile Bold.ttf",
    ("Aptos", False): "Aptos.ttf",
    ("Aptos", True): "Aptos-Bold.ttf",
    ("Aptos Light", False): "Aptos-Light.ttf",
    ("Aptos SemiBold", False): "Aptos-SemiBold.ttf",
    ("Aptos ExtraBold", False): "Aptos-ExtraBold.ttf",
}
_fcache = {}


def style(font="Aptos", size=12, bold=False, color=WHITE, spc=0):
    """spc = letter spacing in points (PowerPoint 'character spacing')."""
    return dict(font=font, size=size, bold=bold, color=color, spc=spc)


def text_mm(text, st):
    key = (st["font"], st["bold"])
    if key not in _fcache:
        _fcache[key] = ImageFont.truetype(os.path.join(FONTDIR, FONT_FILES[key]), 1000)
    w = _fcache[key].getlength(text) / 1000 * st["size"]          # points
    w += st["spc"] * len(text)
    return w * 25.4 / 72


def wrap(text, st, width_mm):
    words, lines, cur = text.split(" "), [], ""
    for w in words:
        trial = (cur + " " + w).strip()
        if cur and text_mm(trial, st) > width_mm:
            lines.append(cur)
            cur = w
        else:
            cur = trial
    lines.append(cur)
    return lines


_cmaps = {}


def check_glyphs(text, st):
    """Fail the build if a character is missing from the font (PowerPoint would
    silently substitute another face, e.g. Eurostile has no middle dot)."""
    from fontTools.ttLib import TTFont
    key = (st["font"], st["bold"])
    if key not in _cmaps:
        _cmaps[key] = TTFont(os.path.join(FONTDIR, FONT_FILES[key])).getBestCmap()
    missing = {c for c in text if ord(c) not in _cmaps[key] and c not in "\n"}
    assert not missing, "%r missing %s in %s" % (text, missing, key)


def mm(v):
    return Emu(int(round(v * 36000)))


# ------------------------------------------------------------------ pptx helpers
class Poster:
    def __init__(self, name):
        self.name = name
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = mm(PAGE_W), mm(PAGE_H)
        self.slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.shapes = self.slide.shapes

    # -- shapes
    def rect(self, x, y, w, h, fill, name, line=None, line_w=0.0):
        s = self.shapes.add_shape(MSO_SHAPE.RECTANGLE, mm(x), mm(y), mm(w), mm(h))
        s.name = name
        if fill:
            s.fill.solid()
            s.fill.fore_color.rgb = RGBColor.from_string(fill)
        else:
            s.fill.background()
        if line:
            s.line.color.rgb = RGBColor.from_string(line)
            s.line.width = Pt(line_w)
        else:
            s.line.fill.background()
        s.shadow.inherit = False
        return s

    def rule(self, x, y, w, color, weight=0.75, name="Rule"):
        s = self.shapes.add_connector(1, mm(x), mm(y), mm(x + w), mm(y))
        s.name = name
        s.line.color.rgb = RGBColor.from_string(color)
        s.line.width = Pt(weight)
        return s

    def picture(self, path, x, y, w=None, h=None, name=None):
        p = self.shapes.add_picture(path, mm(x), mm(y), mm(w) if w else None, mm(h) if h else None)
        if name:
            p.name = name
        return p

    def svg_picture(self, svg_path, png_fallback, x, y, w, h, name):
        """PNG fallback + native SVG (PowerPoint 2016+ renders the SVG as vector)."""
        pic = self.picture(png_fallback, x, y, w, h, name)
        part = Part(PackURI("/ppt/media/%s.svg" % name.replace(" ", "_")),
                    "image/svg+xml", self.prs.part.package,
                    open(svg_path, "rb").read())
        rid = self.slide.part.relate_to(part, RT.IMAGE)
        blip = pic._element.find(".//" + qn("a:blip"))
        ext_lst = etree.SubElement(blip, qn("a:extLst"))
        ext = etree.SubElement(ext_lst, qn("a:ext"))
        ext.set("uri", "{96DAC541-7B7A-43D3-8B79-37D633B846F1}")
        svg = etree.SubElement(
            ext, "{http://schemas.microsoft.com/office/drawing/2016/SVG/main}svgBlip",
            nsmap={"asvg": "http://schemas.microsoft.com/office/drawing/2016/SVG/main"})
        svg.set(qn("r:embed"), rid)
        return pic

    def fit_picture(self, path, cx, cy, max_w, max_h, name):
        from PIL import Image
        iw, ih = Image.open(path).size
        ar = iw / ih
        w = min(max_w, max_h * ar)
        h = w / ar
        return self.picture(path, cx - w / 2, cy - h / 2, w, h, name)

    # -- text
    def text(self, x, y, w, h, paras, name, anchor="t", wrap_text=True):
        """paras: list of (runs, opts); runs: list of (text, style);
        opts: align ('l','c','r'), line (pt, exact), before (pt)."""
        tb = self.shapes.add_textbox(mm(x), mm(y), mm(w), mm(h))
        tb.name = name
        tf = tb.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.word_wrap = wrap_text
        tf.auto_size = None
        tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE,
                              "b": MSO_ANCHOR.BOTTOM}[anchor]
        bodyPr = tf._txBody.find(qn("a:bodyPr"))
        for child in list(bodyPr):
            bodyPr.remove(child)
        etree.SubElement(bodyPr, qn("a:noAutofit"))
        for i, (runs, opts) in enumerate(paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER,
                           "r": PP_ALIGN.RIGHT}[opts.get("align", "l")]
            if "line" in opts:
                p.line_spacing = Pt(opts["line"])
            p.space_before = Pt(opts.get("before", 0))
            p.space_after = Pt(0)
            for t, st in runs:
                for k, piece in enumerate(t.split("\n")):
                    if k:
                        p.add_line_break()
                    self._run(p, piece, st)
            # paragraph mark in the same face, so no theme font (Calibri) is pulled in
            last = p.runs[-1]._r.get_or_add_rPr()
            end = copy.deepcopy(last)
            end.tag = qn("a:endParaRPr")
            p._p.append(end)
        return tb

    @staticmethod
    def _run(p, text, st):
        check_glyphs(text, st)
        r = p.add_run()
        r.text = text
        f = r.font
        f.name = st["font"]
        f.size = Pt(st["size"])
        f.bold = st["bold"]
        f.color.rgb = RGBColor.from_string(st["color"])
        rPr = r._r.get_or_add_rPr()
        if st["spc"]:
            rPr.set("spc", str(int(round(st["spc"] * 100))))
        # make East-Asian / complex-script fallbacks use the same face
        for tag in ("a:ea", "a:cs"):
            el = rPr.find(qn(tag))
            if el is None:
                el = etree.SubElement(rPr, qn(tag))
            el.set("typeface", st["font"])

    def line(self, x, y, w, h, text, st, name, align="l", line=None, anchor="t", wrap_text=True):
        opts = {"align": align}
        if line:
            opts["line"] = line
        return self.text(x, y, w, h, [([(text, st)], opts)], name, anchor, wrap_text)

    def save(self):
        path = os.path.join(OUTDIR, self.name + "_editable.pptx")
        self.prs.core_properties.title = self.name
        self.prs.core_properties.author = "OILS26 organising team"
        self.prs.save(path)
        return path


# ------------------------------------------------------------------ shared pieces
SPONSORS = [  # (file, display name, optical scale) — order as in the approved Roche version
    ("logo_wesemann.png", "Wesemann", 1.08),
    ("logo_roche.png", "Roche", 0.86),
    ("logo_kanton_zurich.png", "Kanton Zürich", 1.38),
    ("logo_microsynth.png", "Microsynth", 0.95),
    ("logo_nucleate.png", "Nucleate", 0.8),
    ("logo_lsz_business.png", "Life Science Zurich Business Network", 1.2),
    ("logo_innovation_zurich.png", "Innovation Zurich", 0.98),
    ("logo_eth_alumni.png", "ETH Alumni", 0.84),
]
QR_URL_SHORT = "b2match.com/e/oils2026"
BAND_Y = 352.0                # top of the white sponsor band
SOCIAL_H = 11.0


def background(P, art):
    P.rect(0, 0, PAGE_W, PAGE_H, NAVY_DEEP, "Background colour")
    P.picture(os.path.join(ASSETS, art), 0, 0, PAGE_W, PAGE_H, "Background artwork")


def logo(P, x, y, w):
    h = w * 645.69 / 1441.85
    P.svg_picture(os.path.join(ASSETS, "oils_logo.svg"),
                  os.path.join(ASSETS, "oils_logo_fallback.png"), x, y, w, h, "OILS logo")
    return h


def qr_block(P, x, y, size, caption, sub):
    pad = size * 0.115
    P.rect(x, y, size, size, WHITE, "QR tile")
    P.picture(os.path.join(ASSETS, "qr_oils2026.png"), x + pad, y + pad,
              size - 2 * pad, size - 2 * pad, "QR code (b2match.com/e/oils2026)")
    cy = y + size + 4.2
    P.line(x - 15, cy, size + 30, 7, caption,
           style("Eurostile", 13.5, True, WHITE, spc=2.2), "QR call to action", align="c")
    P.line(x - 15, cy + 6.6, size + 30, 6, sub,
           style("Aptos", 12, False, SOFT), "QR fallback URL", align="c")


def footer(P):
    band_h = PAGE_H - SOCIAL_H - BAND_Y
    P.rect(0, BAND_Y, PAGE_W, band_h, WHITE, "Sponsor band")
    P.line(M, BAND_Y + 5.2, 120, 6, "SPONSORS & PARTNERS",
           style("Eurostile", 10.5, True, NAVY, spc=3.2), "Sponsors label")
    col = CW / 4
    rows = [BAND_Y + 21.5, BAND_Y + 40.0]
    area = 360.0                                   # target ink area per logo, mm^2
    from PIL import Image
    for i, (f, nm, k) in enumerate(SPONSORS):
        path = os.path.join(ASSETS, f)
        iw, ih = Image.open(path).size
        ar = iw / ih
        w = (area * ar) ** 0.5 * k
        h = w / ar
        if h > 14:
            h = 14; w = h * ar
        if w > 47:
            w = 47; h = w / ar
        cx = M + col * (i % 4) + col / 2
        P.picture(path, cx - w / 2, rows[i // 4] - h / 2, w, h, "Logo " + nm)

    y = PAGE_H - SOCIAL_H
    P.rect(0, y, PAGE_W, SOCIAL_H, NAVY, "Social bar")
    st = style("Aptos SemiBold", 12, False, WHITE)
    P.line(M, y, 40, SOCIAL_H, "Follow us", st, "Follow us", anchor="m")
    ix = M + text_mm("Follow us", st) + 4
    for n in ("linkedin", "facebook", "x", "bluesky"):
        P.picture(os.path.join(ASSETS, "icon_%s.png" % n), ix, y + (SOCIAL_H - 4.6) / 2,
                  4.6, 4.6, "Icon " + n)
        ix += 7.0
    P.line(ix + 1, y, 60, SOCIAL_H, "@OpenLifescience", style("Aptos", 12, False, WHITE),
           "Social handle", anchor="m")
    P.line(PAGE_W - M - 90, y, 90, SOCIAL_H, QR_URL_SHORT,
           style("Aptos SemiBold", 12, False, WHITE), "Event URL", align="r", anchor="m")


# ------------------------------------------------------------------ TEASER
def teaser():
    P = Poster("OILS26_teaser_A3")
    background(P, "bg_teaser.jpg")
    logo(P, M, M, 74)

    P.line(M, 82, CW, 10, "#OILS26 CONFERENCE",
           style("Eurostile", 20, True, SKY, spc=5), "Kicker")

    T = 172                                           # title size, pt
    tb = P.text(M - 2.4, 90, CW + 10, 112,
                [([("From Bit", style("Eurostile", T, True, WHITE, spc=-2))], {"line": T * 0.93}),
                 ([("to Bench", style("Eurostile", T, True, WHITE, spc=-2))], {"line": T * 0.93})],
                "Title", wrap_text=False)
    # optical left alignment: the round 't' sits 2.4 mm further in than the 'F' stem
    tb.text_frame.paragraphs[1]._p.get_or_add_pPr().set("marL", str(int(2.4 * 36000)))

    P.text(M, 202, 200, 26,
           [([("Shaping sustAInable & ethical", style("Aptos Light", 27, False, SOFT))], {"line": 33}),
            ([("innovation in life sciences", style("Aptos Light", 27, False, SOFT))], {"line": 33})],
           "Subtitle")

    P.rule(M, 244, CW, RULE, 0.75, "Divider")

    # date / venue
    y = 253
    P.line(M, y, 190, 22, "20 OCTOBER 2026",
           style("Eurostile", 50, True, SKY, spc=1), "Date", wrap_text=False)
    P.line(M, y + 20.5, 190, 10, "Tuesday  ·  University of Zurich",
           style("Aptos SemiBold", 21, False, WHITE), "Venue")
    P.line(M, y + 29.5, 190, 10, "Rämistrasse 59, Zurich  ·  09:00 – 19:00",
           style("Aptos", 16, False, SOFT), "Venue detail")

    # Roche visit: secondary, typographic only
    ry = 299
    P.rule(M, ry, 166, RULE_SOFT, 0.6, "Roche divider")
    P.line(M, ry + 5, 170, 6, "ALSO ON WEDNESDAY 21 OCTOBER",
           style("Eurostile", 11.5, True, SKY, spc=2.6), "Roche kicker")
    P.line(M, ry + 11, 170, 12, "Site visit to Roche in Basel",
           style("Eurostile", 24, True, WHITE), "Roche headline")
    P.text(M, ry + 23, 172, 14, [
        ([("Limited places for registered OILS26 attendees.", style("Aptos", 14, False, SOFT))], {"line": 18}),
        ([("Express interest by 15 October on the event page.", style("Aptos", 14, False, SOFT))], {"line": 18}),
    ], "Roche detail")

    qs = 60
    qr_block(P, PAGE_W - M - qs, y + 1, qs, "SCAN TO REGISTER", QR_URL_SHORT)
    footer(P)
    return P


# ------------------------------------------------------------------ AGENDA
# (start, end, kind, speakers, title) — titles from the public b2match agenda,
# speakers from the approved Roche-version poster
AGENDA = [
    ("09:00", "", "", "", "Opening & Introduction"),
    ("09:30", "10:15", "KEYNOTE 1", "Lizbé Koekemoer",
     "Feeding the Models: Open Data for AI-Driven Drug Discovery"),
    ("10:15", "11:15", "PANEL", "Jilles Vreeken  ·  Ellen Carbo  ·  Fergus Imrie",
     "Scientific AI: Data Quality, Current Limitations\n& the Future of Models That Understand Science"),
    ("11:15", "", "BREAK", "", "Coffee break  ·  1:1 matchmaking meetings"),
    ("11:45", "13:00", "WORKSHOP 1", "Sebastian Kahlert  ·  Pernilla Sörme",
     "Sustainable Practice in the Lab"),
    ("13:00", "", "BREAK", "", "Lunch  ·  1:1 matchmaking meetings"),
    ("14:00", "15:15", "WORKSHOP 2", "Joseph Heng",
     "Accessing & Using Open AI Tools,\nPre-trained Models and Datasets for Life Sciences"),
    ("15:15", "", "BREAK", "", "Coffee break  ·  1:1 matchmaking meetings"),
    ("15:45", "16:30", "KEYNOTE 2", "Nicoletta Iacobacci",
     "The Missing Behavioural Layer in AI-Assisted Science"),
    ("16:30", "17:30", "PITCH FEST", "",
     "Emerging Voices: PhD & Postdoc Research Showcase"),
    ("17:30", "", "", "", "Closing Ceremony & Pitch Awards"),
    ("17:45", "", "", "", "Networking Apéro in the RAA Lichthof, until 19:00"),
]


def agenda():
    P = Poster("OILS26_agenda_A3")
    background(P, "bg_agenda.jpg")
    lh = logo(P, M, M, 54)

    y = M + lh + 7
    P.line(M, y, 180, 8, "#OILS26 CONFERENCE",
           style("Eurostile", 14, True, SKY, spc=4), "Kicker")
    P.line(M - 0.8, y + 5.5, 200, 26, "From Bit to Bench",
           style("Eurostile", 60, True, WHITE, spc=-0.5), "Title", wrap_text=False)
    P.text(M, y + 29, 210, 11, [([
        ("20 OCTOBER 2026", style("Eurostile", 23, True, SKY, spc=0.5)),
        ("     Tuesday  ·  University of Zurich", style("Aptos SemiBold", 18, False, WHITE)),
    ], {})], "Date and venue")
    P.line(M, y + 40.5, 200, 8, "Rämistrasse 59, Zurich  ·  all sessions in room RAA-G-01",
           style("Aptos", 13.5, False, SOFT), "Room")

    qs = 46
    qr_block(P, PAGE_W - M - qs, M, qs, "SCAN TO REGISTER", "Agenda & updates online")

    # programme header
    top = y + 53
    P.line(M, top, 120, 8, "PROGRAMME", style("Eurostile", 13, True, SKY, spc=4), "Programme label")
    P.line(PAGE_W - M - 160, top + 0.5, 160, 8,
           "Info desk and early-bird 1:1 meetings from 08:00",
           style("Aptos", 12.5, False, MUTED), "Programme note", align="r")
    y = top + 9.5
    P.rule(M, y, CW, RULE, 0.75, "Agenda rule top")

    tx = M + 35                          # text column
    tw = PAGE_W - M - tx
    s_time = style("Eurostile", 22, True, SKY)
    s_time_b = style("Eurostile", 17, True, MUTED)
    s_end = style("Aptos", 12, False, MUTED)
    s_kind = style("Eurostile", 12, True, SKY, spc=2.4)
    s_spk = style("Aptos", 14, False, SOFT)
    s_title = style("Aptos SemiBold", 19, False, WHITE)
    s_brk = style("Aptos", 16, False, SOFT)
    PT = 25.4 / 72
    L_title = 23.0                        # pt line pitch for titles
    L_lab = 6.8                           # mm from label top to title top

    # pass 1: content height of each row
    rows = []
    for start, end, kind, spk, title in AGENDA:
        if kind == "BREAK":
            rows.append(("break", 6.6, None))
        else:
            lines = [l for part in title.split("\n") for l in wrap(part, s_title, tw - 1)]
            h = len(lines) * L_title * PT
            if kind or spk:
                h += L_lab
            rows.append(("talk", max(h, 12.4 if end else 0), lines))

    # the agenda fills the height down to the sponsor band
    avail = BAND_Y - 12 - y
    pad = (avail - sum(r[1] for r in rows)) / (2 * len(rows))
    print("agenda: row padding %.2f mm" % pad)
    assert pad > 1.8, "agenda does not fit"

    for (start, end, kind, spk, title), (typ, h, lines) in zip(AGENDA, rows):
        y0 = y + pad
        if typ == "break":
            P.line(M, y0 - 0.2, 30, 8, start, s_time_b, "Time " + start)
            P.line(tx, y0, tw, 8, title, s_brk, "Break " + start)
        else:
            P.line(M, y0 - 0.6, 32, 10, start, s_time, "Time " + start)
            if end:
                P.line(M, y0 + 7.6, 32, 6, "to " + end, s_end, "End " + start)
            cy = y0
            if kind or spk:
                runs = []
                if kind:
                    runs.append((kind, s_kind))
                if spk:
                    runs.append((("    " if kind else "") + spk, s_spk))
                P.text(tx, cy, tw, 6, [(runs, {})], "Label " + start)
                cy += L_lab
            P.line(tx, cy - 0.6, tw, len(lines) * L_title * PT + 2, title, s_title,
                   "Title " + start, line=L_title)
        y = y0 + h + pad
        P.rule(M, y, CW, RULE_SOFT, 0.5, "Agenda rule")

    footer(P)
    return P


# ------------------------------------------------------------------ export
APPLESCRIPT = '''
on run argv
  set inF to item 1 of argv
  set outF to item 2 of argv
  tell application "Microsoft PowerPoint"
    open (POSIX file inF)
    delay 2
    set pres to active presentation
    save pres in (POSIX file outF) as save as PDF
    close pres saving no
  end tell
end run
'''


def export(pptx):
    pdf = pptx.replace("_editable.pptx", "_print.pdf")
    scr = os.path.join(HERE, ".export.applescript")
    open(scr, "w").write(APPLESCRIPT)
    subprocess.run(["osascript", scr, pptx, pdf], check=True)
    os.remove(scr)
    os.makedirs(PREVIEW, exist_ok=True)
    base = os.path.join(PREVIEW, os.path.basename(pdf).replace("_print.pdf", ""))
    subprocess.run(["pdftoppm", "-png", "-r", "150", "-singlefile", pdf, base + "_preview_150dpi"], check=True)
    subprocess.run(["pdftoppm", "-png", "-scale-to", "700", "-singlefile", pdf, base + "_preview_small"], check=True)
    return pdf


if __name__ == "__main__":
    which = sys.argv[1:] or ["teaser", "agenda"]
    for w in which:
        P = {"teaser": teaser, "agenda": agenda}[w]()
        path = P.save()
        print("pptx", path)
        print("pdf ", export(path))
