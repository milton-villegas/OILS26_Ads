"""A3 versions of the approved posters in 05_Versiones_anteriores/2026-10_Roche_aprobado/.

The originals (285.75 mm wide) are not A3. This script copies them, never touching
the source files, and:

  1. scales every shape, font size, character/line spacing, inset and line weight
     uniformly by 297 / 285.75, so nothing is distorted;
  2. adds the extra page height (A3 is taller in proportion) as whitespace between
     existing blocks: evenly between the agenda rows on the programme, and between
     the teaser's blocks plus a new Roche site-visit panel on the teaser;
  3. swaps the bloom artwork for the same artwork redrawn at A3 / 300 dpi
     (bloom_a3.py), because stretching the old picture would distort the icons.

Then PowerPoint exports each PPTX to PDF and pdftoppm renders the previews.

    python3.12 bloom_a3.py       # artwork (once)
    python3.12 a3_from_roche.py
"""
import copy
import os
import subprocess

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

from build_posters import Poster, check_glyphs, style  # text helpers + glyph check

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, "..", ".."))
SRC = os.path.join(ROOT, "05_Versiones_anteriores", "2026-10_Roche_aprobado")
ASSETS = os.path.join(HERE, "assets")
# outputs go to the archive folder, never to 01_FINAL (which holds hand-edited files)
OUTDIR = os.path.join(ROOT, "05_Versiones_anteriores", "2026-10_A3_generado")
PREVIEW = os.path.join(OUTDIR, "previews")

EMU = 36000
A3_W, A3_H = 297.0, 420.0
S = A3_W / 285.75                       # uniform scale, 1.0394


# ------------------------------------------------------------------ scaling
def scale_text_props(root):
    for el in root.iter():
        tag = etree.QName(el).localname
        if tag in ("rPr", "defRPr", "endParaRPr") and el.get("sz"):
            el.set("sz", str(int(round(int(el.get("sz")) * S))))
        if tag in ("rPr", "defRPr", "endParaRPr") and el.get("spc"):
            el.set("spc", str(int(round(int(el.get("spc")) * S))))
        if tag == "spcPts":
            el.set("val", str(int(round(int(el.get("val")) * S))))
        if tag == "bodyPr":
            for a in ("lIns", "tIns", "rIns", "bIns"):
                if el.get(a):
                    el.set(a, str(int(round(int(el.get(a)) * S))))
        if tag == "ln" and el.get("w"):
            el.set("w", str(int(round(int(el.get("w")) * S))))
        if tag == "pPr":
            for a in ("marL", "indent"):
                if el.get(a):
                    el.set(a, str(int(round(int(el.get(a)) * S))))


def transform(prs, shift):
    """shift(y_mm_original, shape) -> extra vertical offset in final mm."""
    prs.slide_width, prs.slide_height = Emu(int(A3_W * EMU)), Emu(int(A3_H * EMU))
    slide = prs.slides[0]
    for sh in slide.shapes:
        y0 = sh.top / EMU
        dy = shift(y0, sh)
        sh.left = Emu(int(round(sh.left * S)))
        sh.width = Emu(int(round(sh.width * S)))
        sh.height = Emu(int(round(sh.height * S)))
        sh.top = Emu(int(round(sh.top * S + dy * EMU)))
        scale_text_props(sh._element)


def replace_bloom(prs, png):
    slide = prs.slides[0]
    for sh in slide.shapes:
        if sh.name == "Background":
            sh.left = sh.top = 0
            sh.width, sh.height = prs.slide_width, prs.slide_height
        if sh.name == "Bloom artwork":
            img_part, rid = slide.part.get_or_add_image_part(png)
            sh._element.find(".//" + qn("a:blip")).set(qn("r:embed"), rid)
            sh.left = sh.top = 0
            sh.width, sh.height = prs.slide_width, prs.slide_height


# ------------------------------------------------------------------ programme
def programme():
    prs = Presentation(os.path.join(SRC, "oils26_05_bloom_program.pptx"))
    rules = sorted(sh.top / EMU for sh in prs.slides[0].shapes if sh.name.startswith("Agenda rule"))
    band_top = 303.7                                   # original sponsor band top
    extra = A3_H - prs.slide_height / EMU * S          # 48.75 mm
    bottom_gap = 4.5                                   # extra air above the band
    d = (extra - bottom_gap) / len(rules)              # extra per agenda row

    def shift(y, sh):
        if y >= band_top - 0.5:                        # sponsor band, social bar
            return extra
        if y < rules[0] - 12:                          # header incl. divider
            return 0.0
        if sh.name.startswith("Agenda rule"):
            return (rules.index(y) + 1) * d
        row = sum(1 for r in rules if r < y + 1.5)
        return row * d + d / 2

    transform(prs, shift)
    replace_bloom(prs, os.path.join(ASSETS, "bloom_a3_program.png"))
    print("programme: +%.2f mm per agenda row" % d)
    return prs, "OILS26_agenda_A3"


# ------------------------------------------------------------------ teaser
ROCHE_CARD = "FFFFFF"
ROCHE_BLUE = "0B41CD"           # Roche corporate blue, as in the logo
NAVY_TXT = "0B2554"


def teaser():
    prs = Presentation(os.path.join(SRC, "oils26_05_bloom_teaser.pptx"))
    band_top = 253.5
    extra = A3_H - prs.slide_height / EMU * S          # 101 mm
    A, B = 13.0, 21.0                                  # air added above title / above date

    def shift(y, sh):
        if y >= band_top - 0.5:
            return extra
        if y < 50:                                     # logo
            return 0.0
        if y < 166:                                    # kicker, title, subtitle
            return A
        return A + B                                   # divider, date, QR, button

    transform(prs, shift)
    replace_bloom(prs, os.path.join(ASSETS, "bloom_a3_teaser.png"))

    slide = prs.slides[0]
    by = {sh.name: sh for sh in slide.shapes}
    block_bottom = max((sh.top + sh.height) / EMU for sh in slide.shapes
                       if sh.name in ("QR caption", "Register button") and sh.top / EMU < 300)
    band = by["Sponsor band"].top / EMU
    divider = by["Divider"]
    x0 = divider.left / EMU
    w = divider.width / EMU

    # ---- Roche site-visit panel: white card with the Roche logo
    P = Poster.__new__(Poster)
    P.prs, P.slide, P.shapes, P.name = prs, slide, slide.shapes, "teaser"
    gap_above, gap_below = 13.0, 13.0
    card_y = block_bottom + gap_above
    card_h = band - gap_below - card_y
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(int(x0 * EMU)), Emu(int(card_y * EMU)),
                                  Emu(int(w * EMU)), Emu(int(card_h * EMU)))
    card.name = "Roche visit card"
    card.adjustments[0] = 0.09
    card.fill.solid()
    card.fill.fore_color.rgb = RGBColor.from_string(ROCHE_CARD)
    card.line.fill.background()
    card.shadow.inherit = False
    pad = 9.0
    logo_w = 52.0
    from PIL import Image
    iw, ih = Image.open(os.path.join(ASSETS, "logo_roche.png")).size
    logo_h = logo_w * ih / iw
    P.picture(os.path.join(ASSETS, "logo_roche.png"), x0 + pad + 2, card_y + (card_h - logo_h) / 2,
              logo_w, logo_h, "Roche visit logo")
    sep_x = x0 + pad + 2 + logo_w + 9
    sep = slide.shapes.add_connector(1, Emu(int(sep_x * EMU)), Emu(int((card_y + 8) * EMU)),
                                     Emu(int(sep_x * EMU)), Emu(int((card_y + card_h - 8) * EMU)))
    sep.name = "Roche separator"
    sep.line.color.rgb = RGBColor.from_string("C9D3E6")
    sep.line.width = Pt(0.75)

    tx = sep_x + 9
    tw = x0 + w - pad - tx
    ty = card_y + (card_h - 39.5) / 2
    P.line(tx, ty, tw, 7, "ROCHE SITE VISIT  |  WEDNESDAY 21 OCTOBER",
           style("Eurostile", 13, True, ROCHE_BLUE, spc=1.6), "Roche kicker")
    P.line(tx, ty + 7.5, tw, 13, "Visit Roche in Basel",
           style("Eurostile", 30, True, NAVY_TXT), "Roche headline")
    P.text(tx, ty + 21.5, tw, 18, [
        ([("Limited places for registered OILS26 attendees.", style("Eurostile", 16, False, NAVY_TXT))],
         {"line": 20}),
        ([("Express interest by 15 October on the event page.", style("Eurostile", 16, True, NAVY_TXT))],
         {"line": 20}),
    ], "Roche detail")
    print("teaser: Roche card %.1f-%.1f mm, band at %.1f mm" % (card_y, card_y + card_h, band))
    return prs, "OILS26_teaser_A3"


# ------------------------------------------------------------------ export
APPLESCRIPT = '''
on run argv
  tell application "Microsoft PowerPoint"
    open (POSIX file (item 1 of argv))
    delay 2
    set pres to active presentation
    save pres in (POSIX file (item 2 of argv)) as save as PDF
    close pres saving no
  end tell
end run
'''


def export(prs, name):
    pptx = os.path.join(OUTDIR, name + "_editable.pptx")
    pdf = os.path.join(OUTDIR, name + "_print.pdf")
    prs.save(pptx)
    scr = os.path.join(HERE, ".export.applescript")
    open(scr, "w").write(APPLESCRIPT)
    subprocess.run(["osascript", scr, pptx, pdf], check=True)
    os.remove(scr)
    os.makedirs(PREVIEW, exist_ok=True)
    base = os.path.join(PREVIEW, name)
    subprocess.run(["pdftoppm", "-png", "-r", "150", "-singlefile", pdf, base + "_preview_150dpi"], check=True)
    subprocess.run(["pdftoppm", "-png", "-scale-to", "700", "-singlefile", pdf, base + "_preview_small"], check=True)
    print("wrote", pptx, pdf, sep="\n  ")


if __name__ == "__main__":
    import sys
    which = sys.argv[1:] or ["agenda", "teaser"]
    for w in which:
        export(*{"agenda": programme, "teaser": teaser}[w]())
